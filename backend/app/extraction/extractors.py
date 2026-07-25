from pathlib import Path
from typing import Protocol


class UnsupportedFormatError(RuntimeError):
    pass


class TextExtractor(Protocol):
    def __call__(self, path: Path) -> str: ...


def _clean(value: object) -> str:
    return str(value).replace("\x0b", "\n").replace("\ufffd", "-").strip()


def extract_pptx(path: Path) -> str:
    from pptx import Presentation

    presentation = Presentation(str(path))
    sections: list[str] = []
    for slide_number, slide in enumerate(presentation.slides, 1):
        parts: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    values = [_clean(cell.text) for cell in row.cells]
                    if any(values):
                        parts.append(" | ".join(values))
            elif getattr(shape, "has_text_frame", False):
                text = _clean(shape.text)
                if text:
                    parts.append(text)
        if parts:
            sections.append(
                f"--- SLIDE {slide_number} ---\n" + "\n".join(parts)
            )
    return "\n\n".join(sections)


def extract_docx(path: Path) -> str:
    from docx import Document
    from docx.table import Table

    document = Document(str(path))
    parts: list[str] = []
    table_number = 0
    if hasattr(document, "iter_inner_content"):
        content = document.iter_inner_content()
    else:
        content = [*document.paragraphs, *document.tables]
    for block in content:
        if isinstance(block, Table):
            table_number += 1
            rows = []
            for row in block.rows:
                values = [_clean(cell.text) for cell in row.cells]
                if any(values):
                    rows.append(" | ".join(values))
            if rows:
                parts.append(
                    f"--- TABLE {table_number} ---\n" + "\n".join(rows)
                )
        else:
            text = _clean(block.text)
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def extract_pdf(path: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    sections: list[str] = []
    for page_number, page in enumerate(reader.pages, 1):
        text = _clean(page.extract_text() or "")
        if text:
            sections.append(f"--- PAGE {page_number} ---\n{text}")
    return "\n\n".join(sections)


def extract_xlsx(path: Path) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
        keep_links=False,
    )
    sections: list[str] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), 1):
                values = ["" if cell is None else _clean(cell) for cell in row]
                if any(value != "" for value in values):
                    rows.append(f"ROW {row_number}: " + " | ".join(values))
            if rows:
                sections.append(
                    f"--- SHEET: {_clean(sheet.title)} ---\n" + "\n".join(rows)
                )
    finally:
        workbook.close()
    return "\n\n".join(sections)


def extract_legacy_doc(path: Path) -> str:
    raise UnsupportedFormatError(
        "Legacy .doc files require conversion to .docx before extraction."
    )


EXTRACTORS: dict[str, TextExtractor] = {
    ".pptx": extract_pptx,
    ".docx": extract_docx,
    ".pdf": extract_pdf,
    ".xlsx": extract_xlsx,
    ".doc": extract_legacy_doc,
}
