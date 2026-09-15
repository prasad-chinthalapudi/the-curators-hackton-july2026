import json

from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from reportlab.pdfgen import canvas

from app.extraction.models import ExtractionOptions
from app.extraction.pipeline import DocumentExtractionPipeline


def _create_pptx(path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Demand Forecasting"
    table = slide.shapes.add_table(2, 2, 0, 0, 4_000_000, 1_000_000).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Accuracy"
    table.cell(1, 1).text = "95%"
    presentation.save(path)


def _create_docx(path):
    document = Document()
    document.add_paragraph("Healthcare modernization overview")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Technology"
    table.cell(0, 1).text = "Azure"
    table.cell(1, 0).text = "Platform"
    table.cell(1, 1).text = "Databricks"
    document.save(path)


def _create_pdf(path):
    pdf = canvas.Canvas(str(path))
    pdf.drawString(72, 720, "Project outcome: reporting improved by 35 percent.")
    pdf.showPage()
    pdf.save()


def _create_xlsx(path):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Metrics"
    sheet.append(["Name", "Value", "Enabled"])
    sheet.append(["Savings", 0, False])
    workbook.save(path)


def test_pipeline_extracts_all_modern_formats_and_preserves_boundaries(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    _create_pptx(source / "project.pptx")
    _create_docx(source / "project.docx")
    _create_pdf(source / "project.pdf")
    _create_xlsx(source / "project.xlsx")
    output = tmp_path / "output" / "document_index.json"

    index, written = DocumentExtractionPipeline(ExtractionOptions(
        source_directory=source,
        output_file=output,
    )).run()

    assert written == output.resolve()
    assert index.metadata.total_documents == 4
    assert index.metadata.total_errors == 0
    by_type = {document.file_type: document for document in index.documents}
    assert "--- SLIDE 1 ---" in by_type[".pptx"].raw_text
    assert "Metric | Value" in by_type[".pptx"].raw_text
    assert "--- TABLE 1 ---" in by_type[".docx"].raw_text
    assert "--- PAGE 1 ---" in by_type[".pdf"].raw_text
    assert "--- SHEET: Metrics ---" in by_type[".xlsx"].raw_text
    assert "Savings | 0 | False" in by_type[".xlsx"].raw_text


def test_pipeline_reports_legacy_doc_as_structured_error(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    (source / "legacy.doc").write_bytes(b"legacy binary placeholder")
    output = tmp_path / "index.json"
    index, _ = DocumentExtractionPipeline(ExtractionOptions(
        source_directory=source,
        output_file=output,
    )).run()
    assert index.metadata.total_documents == 0
    assert index.metadata.total_errors == 1
    assert index.errors[0].error_type == "UnsupportedFormatError"
    assert "conversion to .docx" in index.errors[0].message


def test_pipeline_is_deterministic_and_honors_limit(tmp_path):
    source = tmp_path / "source"
    source.mkdir()
    for name in ["c.docx", "a.docx", "b.docx"]:
        document = Document()
        document.add_paragraph(name)
        document.save(source / name)
    output = tmp_path / "index.json"
    options = ExtractionOptions(
        source_directory=source,
        output_file=output,
        max_files=2,
    )
    first, _ = DocumentExtractionPipeline(options).run()
    second, _ = DocumentExtractionPipeline(options).run()
    assert [item.filename for item in first.documents] == ["a.docx", "b.docx"]
    assert [item.document_key for item in first.documents] == [
        item.document_key for item in second.documents
    ]
    payload = json.loads(output.read_text(encoding="utf-8"))
    assert payload["metadata"]["total_discovered"] == 2
